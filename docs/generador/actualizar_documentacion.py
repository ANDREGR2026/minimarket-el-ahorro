"""Actualiza los documentos funcionales con los cambios incorporados al sistema."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]

UPDATES = {
    "01_Documento_de_Requerimientos.docx": [
        ("3.9 Vista pública y catálogo consultable", [
            "RF-29. El sistema debe ofrecer una portada pública con información comercial, horario, teléfono, dirección y accesos directos al catálogo y al contacto.",
            "RF-30. La vista pública debe mostrar categorías activas y productos disponibles obtenidos desde el catálogo del sistema.",
            "RF-31. El visitante debe poder buscar productos por nombre o código de barras y filtrar el catálogo por categoría sin iniciar sesión.",
            "RF-32. La información comercial publicada debe obtenerse de la configuración del negocio para evitar datos duplicados en las vistas públicas.",
        ]),
        ("3.10 Operación y continuidad", [
            "RF-33. El usuario debe poder solicitar un enlace de recuperación de contraseña mediante el correo asociado a su cuenta.",
            "RF-34. El administrador debe poder configurar y generar respaldos de la base de datos, así como consultar los archivos generados.",
            "RF-35. El rol Almacenero debe consultar productos, inventario y kardex sin acceder a funciones administrativas o de venta.",
        ]),
    ],
    "02_Historias_de_Usuario.docx": [
        ("HU-19 Consultar el catálogo público", [
            "Como visitante, quiero explorar y buscar los productos publicados para conocer su precio y disponibilidad antes de visitar la tienda.",
            "Criterios de aceptación: el catálogo muestra productos activos, permite buscar por nombre o código y filtrar por categoría; no requiere iniciar sesión.",
        ]),
        ("HU-20 Recuperar el acceso", [
            "Como usuario registrado, quiero solicitar un enlace de recuperación para restablecer mi contraseña cuando no pueda iniciar sesión.",
            "Criterios de aceptación: el sistema valida el correo, no revela si la cuenta existe y permite establecer una contraseña nueva mediante un token vigente.",
        ]),
        ("HU-21 Proteger la continuidad del negocio", [
            "Como administrador, quiero programar y generar respaldos para poder recuperar la información ante una falla.",
            "Criterios de aceptación: se configura la frecuencia y retención, se genera una copia manual y se visualizan los respaldos disponibles.",
        ]),
    ],
    "03_Casos_de_Uso.docx": [
        ("CU-15 Consultar catálogo público", [
            "Actor principal: Visitante.",
            "Flujo principal: el visitante abre la tienda, ingresa un término de búsqueda o elige una categoría y el sistema presenta los productos activos con su nombre, categoría, precio y unidad de medida.",
            "Resultado: el visitante consulta la información comercial sin autenticación y puede continuar hacia los datos de contacto.",
        ]),
        ("CU-16 Recuperar contraseña", [
            "Actor principal: Usuario registrado.",
            "Flujo principal: el usuario solicita la recuperación, recibe un enlace temporal y registra una contraseña nueva. El sistema invalida el token después de utilizarlo.",
        ]),
        ("CU-17 Gestionar respaldos", [
            "Actor principal: Administrador.",
            "Flujo principal: el administrador configura la frecuencia y retención, genera un respaldo cuando lo requiere y descarga los archivos autorizados.",
        ]),
    ],
    "04_Diagramas_UML.docx": [
        ("9 Actualización de alcance", [
            "El modelo actualizado incorpora al Visitante como actor de la vista pública. Este actor consulta el catálogo, filtra por categoría, busca productos y revisa los datos de contacto sin iniciar sesión.",
            "También se incorporan los flujos de recuperación de contraseña y de gestión de respaldos. El Almacenero participa en la consulta de productos, inventario y kardex conforme a los permisos definidos.",
            "Los diagramas fuente en docs/diagramas mantienen la representación técnica de casos de uso, clases, datos, secuencia, actividades y despliegue.",
        ]),
    ],
    "05_Diccionario_de_Datos.docx": [
        ("6 Actualización de configuración y seguridad", [
            "La tabla configuracion centraliza el nombre comercial, moneda, dirección, teléfono, frecuencia de respaldo y cantidad de copias por conservar. Las vistas públicas consumen estos valores para mostrar información vigente del negocio.",
            "La recuperación de contraseña utiliza información temporal asociada a la cuenta del usuario. El token se valida en el servidor, tiene vigencia limitada y deja de ser válido después de restablecer la contraseña.",
            "El catálogo público consulta productos con estado activo y su categoría relacionada. No escribe información ni expone datos de cuentas internas.",
        ]),
    ],
    "06_Manual_de_Usuario.docx": [
        ("13 Sitio público", [
            "La portada pública muestra el horario, teléfono, dirección y accesos al catálogo. Desde Productos puede escribir un nombre o código en el buscador, seleccionar una categoría y consultar precio y unidad de medida sin iniciar sesión.",
            "Las secciones Nosotros, Galería y Contacto presentan información comercial del minimarket. Los datos de contacto se toman de la configuración del sistema.",
        ]),
        ("14 Recuperación de contraseña", [
            "En la pantalla de acceso seleccione Recuperar contraseña, ingrese el correo de su cuenta y revise el mensaje recibido. Abra el enlace antes de su vencimiento y registre la nueva contraseña.",
        ]),
        ("15 Respaldos", [
            "El Administrador puede ingresar a Respaldos para fijar la frecuencia, definir cuántas copias conservar, generar un respaldo manual y descargar un archivo autorizado.",
        ]),
    ],
    "07_Analisis_y_Diseno_de_Arquitectura.docx": [
        ("15 Actualización de arquitectura", [
            "La aplicación incorpora una capa pública separada de las pantallas internas. Las vistas index, nosotros, tienda, galeria y contacto reutilizan componentes de cabecera, navegación y pie, y consumen la configuración comercial mediante el modelo Configuracion.",
            "La vista tienda consulta Producto y Categoria con filtros de búsqueda y categoría. El servidor mantiene la validación de los filtros y publica únicamente productos activos.",
            "La continuidad incorpora BackupController y la recuperación de acceso usa RecuperacionController y Mailer. Ambos flujos preservan la separación existente entre páginas, controladores, modelos y persistencia PDO.",
            "La autorización reconoce Administrador, Cajero y Almacenero. El middleware centraliza el control por rol antes de ejecutar funciones internas.",
        ]),
    ],
    "CASO_PRACTICO_02_Analisis_y_Maquetacion.docx": [
        ("6 Actualización de la práctica", [
            "La maquetación evolucionó desde un panel interno hacia una experiencia completa con vista pública. La portada prioriza datos operativos, acceso a categorías y productos destacados; el catálogo agrega búsqueda y filtrado por categoría.",
            "Las pantallas públicas reutilizan la paleta y el sistema de componentes del panel administrativo, pero reducen acciones a las que necesita un visitante: conocer el negocio, consultar productos y contactar la tienda.",
            "El material de referencia actualizado se encuentra en figma, organizado por roles, y en docs/capturas para las pantallas implementadas.",
        ]),
    ],
}

def remove_previous_appendix(doc):
    marker = "Actualización de documentación"
    start = None
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == marker:
            start = index
            break
    if start is not None:
        for paragraph in doc.paragraphs[start:]:
            paragraph._element.getparent().remove(paragraph._element)

def append_update(doc, sections):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    title = doc.add_paragraph("Actualización de documentación")
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(16)
    title.paragraph_format.space_before = Pt(0)
    intro = doc.add_paragraph(
        "Esta sección registra los cambios incorporados al Sistema Web para la Gestión de un Minimarket después de la versión documentada inicialmente."
    )
    intro.paragraph_format.space_after = Pt(10)
    for heading, items in sections:
        subheading = doc.add_paragraph(heading)
        subheading.runs[0].bold = True
        subheading.runs[0].font.size = Pt(13)
        for item in items:
            paragraph = doc.add_paragraph("• " + item)
            paragraph.paragraph_format.left_indent = Pt(18)
            paragraph.paragraph_format.first_line_indent = Pt(-10)

def main():
    for filename, sections in UPDATES.items():
        path = ROOT / filename
        doc = Document(path)
        remove_previous_appendix(doc)
        append_update(doc, sections)
        doc.save(path)
        print(f"Actualizado: {path.name}")

if __name__ == "__main__":
    main()
